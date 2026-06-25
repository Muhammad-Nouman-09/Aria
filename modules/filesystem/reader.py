"""File reading + directory listing.

Phase 1 reads plain-text/code formats natively and handles .docx/.xlsx/.pdf
when the optional libraries are installed (they ship in the full
requirements.txt). Binary formats degrade to a clear message rather than
dumping garbage into the model context.
"""

from __future__ import annotations

import csv
import io
from pathlib import Path

TEXT_SUFFIXES = {
    ".txt", ".md", ".py", ".js", ".ts", ".tsx", ".jsx", ".json", ".csv",
    ".tsv", ".html", ".htm", ".xml", ".yaml", ".yml", ".ini", ".cfg",
    ".toml", ".log", ".sql", ".sh", ".ps1", ".bat", ".c", ".cpp", ".h",
    ".java", ".cs", ".go", ".rs", ".rb", ".php", ".css", ".env",
}


def read_file(file_path: str, encoding: str = "utf-8",
              max_chars: int = 50_000) -> dict:
    p = Path(file_path).expanduser()
    if not p.exists():
        return {"error": f"File not found: {file_path}"}
    if p.is_dir():
        return {"error": f"Path is a directory, not a file: {file_path}"}

    suffix = p.suffix.lower()
    try:
        if suffix in TEXT_SUFFIXES or suffix == "":
            text = p.read_text(encoding=encoding, errors="replace")
        elif suffix == ".docx":
            text = _read_docx(p)
        elif suffix == ".xlsx":
            text = _read_xlsx(p)
        elif suffix == ".pdf":
            text = _read_pdf(p)
        else:
            return {"error": f"Unsupported file type for reading: {suffix}"}
    except ModuleNotFoundError as e:
        return {"error": f"Missing library to read {suffix}: {e.name} "
                         f"(install full requirements.txt)."}
    except Exception as e:
        return {"error": f"Failed to read {file_path}: {e}"}

    truncated = len(text) > max_chars
    if truncated:
        text = text[:max_chars] + "\n...[truncated]"
    return {"path": str(p.resolve()), "chars": len(text),
            "truncated": truncated, "content": text}


def _read_docx(p: Path) -> str:
    import docx  # python-docx
    doc = docx.Document(str(p))
    return "\n".join(para.text for para in doc.paragraphs)


def _read_xlsx(p: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(str(p), read_only=True, data_only=True)
    out = io.StringIO()
    for ws in wb.worksheets:
        out.write(f"# Sheet: {ws.title}\n")
        writer = csv.writer(out)
        for row in ws.iter_rows(values_only=True):
            writer.writerow(["" if c is None else c for c in row])
        out.write("\n")
    return out.getvalue()


def _read_pdf(p: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(p))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def list_directory(directory_path: str, recursive: bool = False,
                   include_hidden: bool = False,
                   filter_extension: str | None = None) -> dict:
    d = Path(directory_path).expanduser()
    if not d.exists():
        return {"error": f"Directory not found: {directory_path}"}
    if not d.is_dir():
        return {"error": f"Not a directory: {directory_path}"}

    items = []
    iterator = d.rglob("*") if recursive else d.iterdir()
    for entry in iterator:
        if not include_hidden and entry.name.startswith("."):
            continue
        if filter_extension and entry.is_file() and entry.suffix.lower() != filter_extension.lower():
            continue
        try:
            size = entry.stat().st_size if entry.is_file() else None
        except OSError:
            size = None
        items.append({
            "name": str(entry.relative_to(d)) if recursive else entry.name,
            "type": "dir" if entry.is_dir() else "file",
            "size_bytes": size,
        })
        if len(items) >= 1000:
            break
    items.sort(key=lambda x: (x["type"] != "dir", x["name"].lower()))
    return {"path": str(d.resolve()), "count": len(items), "items": items}


def search_files(search_path: str, pattern: str,
                 search_type: str = "filename", recursive: bool = True) -> dict:
    d = Path(search_path).expanduser()
    if not d.exists():
        return {"error": f"Path not found: {search_path}"}
    matches: list[str] = []

    if search_type == "filename":
        glob = d.rglob(pattern) if recursive else d.glob(pattern)
        for m in glob:
            matches.append(str(m))
            if len(matches) >= 500:
                break
    else:  # content search
        files = d.rglob("*") if recursive else d.iterdir()
        needle = pattern.lower()
        for f in files:
            if not f.is_file() or f.suffix.lower() not in TEXT_SUFFIXES:
                continue
            try:
                if needle in f.read_text(encoding="utf-8", errors="ignore").lower():
                    matches.append(str(f))
            except Exception:
                continue
            if len(matches) >= 200:
                break
    return {"pattern": pattern, "count": len(matches), "matches": matches}
